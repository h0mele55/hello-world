import sys
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment

OUT = sys.argv[1]
N = 50            # project capacity
MAXDUR = 24       # max duration (months)
TL = 36           # timeline months
R0, R1 = 6, 6 + N - 1   # project rows on every project sheet

wb = Workbook()
wb._named_styles['Normal'].font = Font(name='Arial', size=10)

BLUE = Font(name='Arial', size=10, color='0000FF')
GREEN = Font(name='Arial', size=10, color='008000')
BOLD = Font(name='Arial', size=10, bold=True)
HDR = Font(name='Arial', size=10, bold=True, color='FFFFFF')
TITLE = Font(name='Arial', size=14, bold=True)
GREY = Font(name='Arial', size=8, color='808080')
YELLOW = PatternFill('solid', fgColor='FFFF00')
INFILL = PatternFill('solid', fgColor='FFF2CC')   # light input fill for grids
HDRFILL = PatternFill('solid', fgColor='C0392B')
CALCFILL = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY = '#,##0;(#,##0);-'
PCT = '0.0%;(0.0%);-'
MON = 'mmm yyyy'
WRAP = Alignment(wrap_text=True, vertical='center', horizontal='center')

CM = 'Projects!$C$2'   # current month
TS = 'Projects!$C$3'   # timeline start


def header(ws, row, col, labels, widths=None):
    for i, t in enumerate(labels):
        c = ws.cell(row, col + i, t)
        c.font, c.fill, c.alignment, c.border = HDR, HDRFILL, WRAP, BOX
    ws.row_dimensions[row].height = 32
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[L(col + i)].width = w


def title(ws, text, sub=None):
    ws['A1'] = text
    ws['A1'].font = TITLE
    if sub:
        ws['A3'] = sub
        ws['A3'].font = Font(name='Arial', size=9, italic=True, color='595959')


def mi(month, start):
    """1-based month index of `month` relative to `start`."""
    return f'((YEAR({month})-YEAR({start}))*12+MONTH({month})-MONTH({start})+1)'


# ---------------------------------------------------------------- README
rd = wb.active
rd.title = 'README'
rd.column_dimensions['A'].width = 3
rd.column_dimensions['B'].width = 26
rd.column_dimensions['C'].width = 110
rd['B1'] = 'Project Budget Calculator'
rd['B1'].font = TITLE
rows = [
    ('PURPOSE', 'Calculates the Budgeted balance of each project and phases it over the remaining months of the project.'),
    ('', ''),
    ('FORMULAS', ''),
    ('Budgeted', 'Budgeted = Contracted Budget − WIP to date (the sum of WIP logged in months up to and including the Current Month).'),
    ('Phased Budgeted', 'Budgeted is spread only over the project months AFTER the Current Month, using that project\'s phasing weights for those months, renormalised to 1.0.'),
    ('', 'Phased Budgeted (month k) = Budgeted × weight(k) ÷ Σ weights of remaining months.'),
    ('Weights', 'Enter raw weights (e.g. 1,2,2,2,2). They don\'t need to sum to 1. If a project has no weights at all, every month gets an equal weight.'),
    ('Example (from the slide)', 'Contract 10,000. 5 months, weights 1:2:2:2:2. 1,000 of WIP logged in month I (the current month).'),
    ('', '→ Budgeted = 9,000, phased over months II–V (weights 2:2:2:2 → 25% each) = 2,250 per month.'),
    ('', ''),
    ('SHEETS', ''),
    ('Projects', 'Project ∑ table (page 1). Holds the global Current Month field. Add a project by filling in the next empty row. Use the header filter buttons to filter.'),
    ('Phasing', 'Raw phasing weights per project, month 1–24 of the project. Also shows the weights normalised to 1.0 over the full duration.'),
    ('WIP Log', 'Monthly WIP actuals per project on the calendar timeline. Only months ≤ Current Month count towards WIP to date.'),
    ('Project Page', 'Page 2: pick a project from the dropdown to see its details, month-by-month phasing and chart (the "visualize" view).'),
    ('Timeline', 'Page 3: Gantt-style view of every project. Each cell shows the WIP actual (≤ Current Month) or the Phased Budgeted (after it).'),
    ('Aggregator', 'Page 4: Σ WIP + Σ Phased Budgeted per month against the manual TRM input, with a chart and calendar-year totals.'),
    ('Phased Budget', 'Calculation grid (no inputs): Phased Budgeted per project per calendar month.'),
    ('', ''),
    ('HOW TO USE', ''),
    ('Yellow cells', 'Key inputs: Current Month, Timeline Start, the TRM row and the project selector.'),
    ('Blue text', 'Hard-coded inputs: project fields, weights, WIP amounts. Black text = formulas. Green text = links to another sheet.'),
    ('Light yellow grids', 'Input grids on Phasing and WIP Log. Grey cells there are outside the project\'s duration or window.'),
    ('Monthly routine', '1) Log the month\'s WIP on WIP Log.  2) Move Current Month forward one month on Projects.  Everything recalculates.'),
    ('Status column', 'Projects!N flags problems: WIP above contract, no remaining months left to phase the budget into, missing fields, etc.'),
    ('', ''),
    ('LIMITS & ASSUMPTIONS', ''),
    ('Capacity', f'{N} projects, durations up to {MAXDUR} months, calendar timeline of {TL} months from Timeline Start (default Jan 2026 – Dec 2028).'),
    ('Timeline Start', 'Change it only before logging WIP: the WIP Log columns are tied to calendar months, so moving the start shifts those columns.'),
    ('Months', 'Starting Month and Current Month are compared by year and month only (the day is ignored).'),
    ('Units', 'All amounts are in one currency unit (enter 10000 for 10k). No currency symbol is applied.'),
    ('TRM', 'Monthly comparison target, entered manually on the Aggregator sheet (per user).'),
    ('Annual totals', 'Calendar years (Jan–Dec), WIP actuals + Phased Budgeted.'),
    ('Source', 'Logic from "Revenue_Forecast_Calculator.pptx", with the user\'s clarifications from 2026-10-06.'),
]
for i, (a, b) in enumerate(rows, start=3):
    rd.cell(i, 2, a).font = BOLD if a.isupper() or a else BOLD
    rd.cell(i, 3, b).alignment = Alignment(wrap_text=True, vertical='top')
    if a.isupper() and a:
        rd.cell(i, 2).font = Font(name='Arial', size=11, bold=True, color='C0392B')
rd['B23'].fill = YELLOW
rd['B24'].font = Font(name='Arial', size=10, bold=True, color='0000FF')
rd['B25'].fill = INFILL

# ---------------------------------------------------------------- Projects
pj = wb.create_sheet('Projects')
pj['A1'] = 'Project ∑ Table'
pj['A1'].font = TITLE
pj['A2'] = 'Current Month'
pj['A3'] = 'Timeline Start'
for c in ('A2', 'A3'):
    pj[c].font = BOLD
pj['C2'] = date(2026, 10, 1)
pj['C3'] = date(2026, 1, 1)
for c in ('C2', 'C3'):
    pj[c].font, pj[c].fill, pj[c].number_format, pj[c].border = BLUE, YELLOW, MON, BOX
pj['D2'] = 'Months up to and including this month are actuals (WIP). Later months are forecast (Phased Budgeted).'
pj['D3'] = f'First month of the {TL}-month calendar timeline. Change it only before logging WIP.'
for c in ('D2', 'D3'):
    pj[c].font = Font(name='Arial', size=9, italic=True, color='595959')
pj['C2'].comment = Comment('Default: October 2026 (the build date). Move it forward each month.', 'Claude')
pj['C3'].comment = Comment('Default: Jan 2026. The timeline is 36 months long.', 'Claude')

cols = ['#', 'Project Name', 'Engagement Manager', 'Starting Month', 'Duration (months)', 'End Month',
        'Contracted Budget', 'WIP to Date', 'Budgeted (Contract − WIP)', 'Remaining Months',
        'Phased Budgeted – Next Month', 'Phased Budgeted – Rest of Current Year',
        'Current-Year Forecast (WIP + Phased)', 'Status']
header(pj, 5, 1, cols, [5, 28, 22, 13, 11, 13, 15, 15, 16, 11, 15, 17, 17, 34])
PB = "'Phased Budget'"
for r in range(R0, R1 + 1):
    n = r - R0 + 1
    pj.cell(r, 1, n).font = GREY
    for c in (2, 3, 4, 5, 7):
        pj.cell(r, c).font = BLUE
    pj.cell(r, 4).number_format = MON
    pj.cell(r, 6, f'=IF(OR(D{r}="",E{r}=""),"",EDATE(D{r},E{r}-1))').number_format = MON
    pj.cell(r, 7).number_format = MONEY
    pj.cell(r, 8, f"=IF(B{r}=\"\",\"\",'WIP Log'!C{r})").font = GREEN
    pj.cell(r, 9, f'=IF(B{r}="","",N(G{r})-N(H{r}))')
    pj.cell(r, 10, f'=IF(OR(B{r}="",E{r}=""),"",MAX(0,E{r}-Phasing!F{r}+1))')
    # next month = column index of (CM+1) in the timeline
    nxt = f'{mi(CM, TS)}+1'
    pj.cell(r, 11, f'=IF(B{r}="","",IF(AND({nxt}>=1,{nxt}<={TL}),INDEX({PB}!$H{r}:${L(7+TL)}{r},{nxt}),0))')
    pj.cell(r, 12, f'=IF(B{r}="","",SUMPRODUCT({PB}!$H{r}:${L(7+TL)}{r}*(YEAR({PB}!$H$5:${L(7+TL)}$5)=YEAR({CM}))))')
    pj.cell(r, 13, f"=IF(B{r}=\"\",\"\",SUMPRODUCT('WIP Log'!$D{r}:${L(3+TL)}{r}*('WIP Log'!$D$5:${L(3+TL)}$5<={CM})*(YEAR('WIP Log'!$D$5:${L(3+TL)}$5)=YEAR({CM})))+L{r})")
    for c in range(8, 14):
        pj.cell(r, c).number_format = MONEY
    pj.cell(r, 14, (
        f'=IF(B{r}="","",IF(OR(D{r}="",E{r}="",G{r}=""),"Missing start / duration / contract",'
        f'IF(OR(E{r}<1,E{r}>{MAXDUR}),"Duration must be 1–{MAXDUR}",'
        f'IF(I{r}<0,"WIP exceeds contract",'
        f'IF(AND(J{r}=0,I{r}<>0),"Budget not phased: no remaining months",'
        f'IF(AND(J{r}>0,Phasing!G{r}=0,I{r}<>0),"Budget not phased: remaining weights are 0",'
        f'IF(OR(D{r}<{TS},F{r}>EDATE({TS},{TL - 1})),"Partly outside timeline",'
        f'IF(Phasing!H{r}=0,"OK","Weights entered beyond duration (ignored)"))))))))'))
    for c in range(1, 15):
        pj.cell(r, c).border = BOX
        if c in (6, 8, 9, 10, 11, 12, 13, 14):
            pj.cell(r, c).fill = CALCFILL
# Example row (from the slide, adapted per user clarification)
ex = R0
pj.cell(ex, 2, 'Example Project (slide)')
pj.cell(ex, 3, 'J. Doe')
pj.cell(ex, 4, date(2026, 10, 1))
pj.cell(ex, 5, 5)
pj.cell(ex, 7, 10000)
pj.cell(ex, 2).comment = Comment('Example row from the slide: contract 10k, weights 1:2:2:2:2, 1k of WIP in month I. Overwrite or delete it.', 'Claude')
# totals row
tr = R1 + 2
pj.cell(tr, 2, 'TOTAL').font = BOLD
for c in (7, 8, 9, 11, 12, 13):
    cell = pj.cell(tr, c, f'=SUM({L(c)}{R0}:{L(c)}{R1})')
    cell.font, cell.number_format, cell.border = BOLD, MONEY, Border(top=Side(style='medium'))
pj.cell(tr, 14, f'=COUNTIFS(N{R0}:N{R1},"<>OK",B{R0}:B{R1},"<>")&" project(s) need attention"').font = BOLD
pj.auto_filter.ref = f'A5:N{R1}'
pj.freeze_panes = 'C6'
dv = DataValidation(type='whole', operator='between', formula1='1', formula2=str(MAXDUR), allow_blank=True,
                    showErrorMessage=True, errorTitle='Duration', error=f'Enter a whole number of months from 1 to {MAXDUR}.')
pj.add_data_validation(dv)
dv.add(f'E{R0}:E{R1}')
dvd = DataValidation(type='date', operator='greaterThan', formula1='36526', allow_blank=True,
                     showErrorMessage=True, errorTitle='Date', error='Enter a date (the first of the month is recommended).')
pj.add_data_validation(dvd)
dvd.add(f'D{R0}:D{R1}')
dvd.add('C2:C3')
pj.conditional_formatting.add(f'N{R0}:N{R1}', FormulaRule(formula=[f'AND(N{R0}<>"OK",N{R0}<>"")'],
                              font=Font(color='C00000', bold=True), fill=PatternFill('solid', fgColor='FDE9E7')))

# ---------------------------------------------------------------- Phasing
ph = wb.create_sheet('Phasing')
title(ph, 'Phasing Weights', 'Enter raw weights in the yellow cells (month 1 = Starting Month). They are renormalised to 1.0 automatically. Leave a row blank for equal weights.')
W0 = 9  # first weight column (I)
W1 = W0 + MAXDUR - 1
NW0 = W1 + 2
NW1 = NW0 + MAXDUR - 1
header(ph, 5, 1, ['#', 'Project Name', 'Duration', 'Weights entered', 'Σ weights (full duration)',
                  'First remaining month #', 'Σ weights (remaining months)', 'Non-zero weights beyond duration'],
       [5, 28, 9, 10, 12, 12, 13, 12])
header(ph, 5, W0, [f'M{k}' for k in range(1, MAXDUR + 1)], [7] * MAXDUR)
header(ph, 5, NW0, [f'M{k}' for k in range(1, MAXDUR + 1)], [7] * MAXDUR)
ph.cell(4, W0, 'RAW WEIGHTS (input) →').font = BOLD
ph.cell(4, NW0, 'NORMALISED WEIGHTS, FULL DURATION (Σ = 1.0) →').font = BOLD
ph.column_dimensions[L(W1 + 1)].width = 3
ph.row_dimensions[5].height = 46
# row 2 = numeric month index helper
ph.cell(2, W0 - 1, 'month index').font = GREY
for k in range(1, MAXDUR + 1):
    ph.cell(2, W0 + k - 1, k).font = GREY
    ph.cell(2, NW0 + k - 1, k).font = GREY
WR = lambda r: f'${L(W0)}{r}:${L(W1)}{r}'
IDX = f'${L(W0)}$2:${L(W1)}$2'
for r in range(R0, R1 + 1):
    ph.cell(r, 1, r - R0 + 1).font = GREY
    ph.cell(r, 2, f'=IF(Projects!B{r}="","",Projects!B{r})').font = GREEN
    ph.cell(r, 3, f'=IF(Projects!E{r}="","",Projects!E{r})').font = GREEN
    ph.cell(r, 4, f'=COUNT({WR(r)})')
    ph.cell(r, 5, f'=IF(C{r}="",0,IF(D{r}=0,C{r},SUMPRODUCT({WR(r)}*({IDX}<=C{r}))))')
    ph.cell(r, 6, f'=IF(OR(C{r}="",Projects!D{r}=""),1,MAX(1,{mi(CM, f"Projects!D{r}")}+1))')
    ph.cell(r, 7, f'=IF(C{r}="",0,IF(D{r}=0,MAX(0,C{r}-F{r}+1),SUMPRODUCT({WR(r)}*({IDX}>=F{r})*({IDX}<=C{r}))))')
    ph.cell(r, 8, f'=IF(C{r}="",0,SUMPRODUCT(({WR(r)}<>0)*({IDX}>C{r})))')
    for c in range(1, 9):
        ph.cell(r, c).border = BOX
    for c in range(3, 9):
        ph.cell(r, c).fill = CALCFILL
    for k in range(1, MAXDUR + 1):
        c = ph.cell(r, W0 + k - 1)
        c.font, c.border, c.number_format = BLUE, BOX, '0.##'
        w = f'{L(W0 + k - 1)}{r}'
        n = ph.cell(r, NW0 + k - 1,
                    f'=IF(OR(C{r}="",{L(NW0 + k - 1)}$2>C{r},E{r}=0),"",IF(D{r}=0,1/C{r},N({w})/E{r}))')
        n.number_format, n.border, n.fill = PCT, BOX, CALCFILL
    ph.cell(r, NW1 + 1, f'=IF(C{r}="","",SUM({L(NW0)}{r}:{L(NW1)}{r}))').number_format = PCT
ph.cell(5, NW1 + 1, 'Check Σ').font = BOLD
for k, w in enumerate([1, 2, 2, 2, 2]):
    ph.cell(R0, W0 + k, w)
# input fill only within duration
ph.conditional_formatting.add(f'{L(W0)}{R0}:{L(W1)}{R1}',
                              FormulaRule(formula=[f'AND($C{R0}<>"",{L(W0)}$2<=$C{R0})'], fill=INFILL))
ph.conditional_formatting.add(f'{L(W0)}{R0}:{L(W1)}{R1}',
                              FormulaRule(formula=[f'AND($C{R0}<>"",{L(W0)}$2>$C{R0})'], fill=PatternFill('solid', fgColor='D9D9D9')))
dvw = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True,
                     showErrorMessage=True, errorTitle='Weight', error='Weights must be 0 or more.')
ph.add_data_validation(dvw)
dvw.add(f'{L(W0)}{R0}:{L(W1)}{R1}')
ph.freeze_panes = ph.cell(R0, W0)

# ---------------------------------------------------------------- WIP Log
wl = wb.create_sheet('WIP Log')
title(wl, 'WIP Log (monthly actuals)', 'Enter the WIP actual for each project and month. Only months up to and including the Current Month count; later months are ignored.')
M0, M1 = 4, 4 + TL - 1
header(wl, 5, 1, ['#', 'Project Name', 'WIP to Date (≤ Current Month)'], [5, 28, 15])
for j in range(TL):
    c = wl.cell(5, M0 + j, f'={TS}' if j == 0 else f'=EDATE({L(M0 + j - 1)}5,1)')
    c.font, c.fill, c.alignment, c.border, c.number_format = HDR, HDRFILL, WRAP, BOX, 'mmm yy'
    wl.column_dimensions[L(M0 + j)].width = 9
    wl.cell(4, M0 + j, f'=IF({L(M0 + j)}5<={CM},"actual","future")').font = GREY
for r in range(R0, R1 + 1):
    wl.cell(r, 1, r - R0 + 1).font = GREY
    wl.cell(r, 2, f'=IF(Projects!B{r}="","",Projects!B{r})').font = GREEN
    wl.cell(r, 3, f'=SUMPRODUCT(${L(M0)}{r}:${L(M1)}{r}*(${L(M0)}$5:${L(M1)}$5<={CM}))')
    wl.cell(r, 3).number_format, wl.cell(r, 3).fill = MONEY, CALCFILL
    for c in range(1, 4):
        wl.cell(r, c).border = BOX
    for j in range(TL):
        c = wl.cell(r, M0 + j)
        c.font, c.border, c.number_format = BLUE, BOX, MONEY
wl.cell(R0, M0 + 9, 1000)   # Oct 2026 = month I of the example
tot = R1 + 2
wl.cell(tot, 2, 'TOTAL').font = BOLD
for c in range(3, M1 + 1):
    x = wl.cell(tot, c, f'=SUM({L(c)}{R0}:{L(c)}{R1})')
    x.font, x.number_format = BOLD, MONEY
win = (f'IF(OR(Projects!$D{R0}="",Projects!$E{R0}=""),FALSE,AND({mi(f"{L(M0)}$5", f"Projects!$D{R0}")}>=1,'
       f'{mi(f"{L(M0)}$5", f"Projects!$D{R0}")}<=Projects!$E{R0}))')
rng = f'{L(M0)}{R0}:{L(M1)}{R1}'
wl.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({win},{L(M0)}$5<={CM})'], fill=INFILL))
wl.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({L(M0)}{R0}<>0,{L(M0)}$5>{CM})'],
                              fill=PatternFill('solid', fgColor='FDE9E7'), font=Font(color='C00000', strike=True)))
wl.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($B{R0}<>"",NOT({win}))'], fill=PatternFill('solid', fgColor='EDEDED')))
wl.freeze_panes = wl.cell(R0, M0)

# ---------------------------------------------------------------- Phased Budget (calc)
pb = wb.create_sheet('Phased Budget')
title(pb, 'Phased Budgeted – calculation grid (no inputs)',
      'Budgeted × weight(k) ÷ Σ remaining weights, for project months after the Current Month only.')
G0, G1 = 8, 8 + TL - 1
header(pb, 5, 1, ['#', 'Project Name', 'Starting Month', 'Duration', 'Budgeted', 'Σ remaining weights', 'Equal weights?'],
       [5, 28, 11, 9, 13, 12, 9])
for j in range(TL):
    c = pb.cell(5, G0 + j, f"='WIP Log'!{L(M0 + j)}5")
    c.font, c.fill, c.alignment, c.border, c.number_format = HDR, HDRFILL, WRAP, BOX, 'mmm yy'
    pb.column_dimensions[L(G0 + j)].width = 9
for r in range(R0, R1 + 1):
    pb.cell(r, 1, r - R0 + 1).font = GREY
    pb.cell(r, 2, f'=IF(Projects!B{r}="","",Projects!B{r})').font = GREEN
    pb.cell(r, 3, f'=IF(Projects!D{r}="","",Projects!D{r})').font = GREEN
    pb.cell(r, 3).number_format = MON
    pb.cell(r, 4, f'=IF(Projects!E{r}="","",Projects!E{r})').font = GREEN
    pb.cell(r, 5, f'=IF(Projects!B{r}="",0,N(Projects!I{r}))').font = GREEN
    pb.cell(r, 5).number_format = MONEY
    pb.cell(r, 6, f'=Phasing!G{r}').font = GREEN
    pb.cell(r, 7, f'=Phasing!D{r}=0').font = GREEN
    for c in range(1, 8):
        pb.cell(r, c).border = BOX
    for j in range(TL):
        m = f'{L(G0 + j)}$5'
        k = mi(m, f'$C{r}')
        f = (f'=IF(OR($C{r}="",$D{r}="",$F{r}=0),0,IF(OR({m}<={CM},{k}<1,{k}>$D{r}),0,'
             f'$E{r}*IF($G{r},1,N(INDEX(Phasing!${L(W0)}{r}:${L(W1)}{r},{k})))/$F{r}))')
        c = pb.cell(r, G0 + j, f)
        c.number_format, c.border = MONEY, BOX
pb.cell(tot, 2, 'TOTAL').font = BOLD
for c in range(5, G1 + 1):
    if c in (6, 7):
        continue
    x = pb.cell(tot, c, f'=SUM({L(c)}{R0}:{L(c)}{R1})')
    x.font, x.number_format = BOLD, MONEY
pb.freeze_panes = pb.cell(R0, G0)

# ---------------------------------------------------------------- Timeline
tl = wb.create_sheet('Timeline')
title(tl, 'Project Timeline', 'Dark orange = WIP actual (≤ Current Month). Light orange = Phased Budgeted. The bordered column is the Current Month.')
T0, T1 = 6, 6 + TL - 1
header(tl, 5, 1, ['#', 'Project Name', 'Engagement Manager', 'Start', 'End'], [5, 28, 18, 10, 10])
for j in range(TL):
    c = tl.cell(5, T0 + j, f"='WIP Log'!{L(M0 + j)}5")
    c.font, c.fill, c.alignment, c.border, c.number_format = HDR, HDRFILL, WRAP, BOX, 'mmm yy'
    tl.column_dimensions[L(T0 + j)].width = 8
for r in range(R0, R1 + 1):
    tl.cell(r, 1, r - R0 + 1).font = GREY
    tl.cell(r, 2, f'=IF(Projects!B{r}="","",Projects!B{r})').font = GREEN
    tl.cell(r, 3, f'=IF(Projects!C{r}="","",Projects!C{r})').font = GREEN
    tl.cell(r, 4, f'=IF(Projects!D{r}="","",Projects!D{r})').font = GREEN
    tl.cell(r, 5, f'=IF(Projects!F{r}="","",Projects!F{r})').font = GREEN
    tl.cell(r, 4).number_format = tl.cell(r, 5).number_format = 'mmm yy'
    for j in range(TL):
        m = f'{L(T0 + j)}$5'
        wipc = f"'WIP Log'!{L(M0 + j)}{r}"
        pbc = f"'Phased Budget'!{L(G0 + j)}{r}"
        inwin = f'IF($D{r}="",FALSE,AND({m}>=DATE(YEAR($D{r}),MONTH($D{r}),1),{m}<=$E{r}))'
        f = f'=IF($B{r}="","",IF({m}<={CM},IF(OR({inwin},N({wipc})<>0),N({wipc}),""),IF({inwin},{pbc},"")))'
        c = tl.cell(r, T0 + j, f)
        c.number_format, c.border = '#,##0;(#,##0);0', BOX
        c.font = Font(name='Arial', size=8)
trng = f'{L(T0)}{R0}:{L(T1)}{R1}'
tinwin = f'IF($D{R0}="",FALSE,AND({L(T0)}$5>=DATE(YEAR($D{R0}),MONTH($D{R0}),1),{L(T0)}$5<=$E{R0}))'
tl.conditional_formatting.add(trng, FormulaRule(formula=[f'AND({tinwin},{L(T0)}$5<={CM})'],
                              fill=PatternFill('solid', fgColor='E67E22'), font=Font(color='FFFFFF')))
tl.conditional_formatting.add(trng, FormulaRule(formula=[tinwin], fill=PatternFill('solid', fgColor='FAD7A0')))
tl.conditional_formatting.add(f'{L(T0)}5:{L(T1)}{R1}',
                              FormulaRule(formula=[f'AND(YEAR({L(T0)}$5)=YEAR({CM}),MONTH({L(T0)}$5)=MONTH({CM}))'],
                                          border=Border(left=Side(style='medium', color='C0392B'), right=Side(style='medium', color='C0392B'))))
tl.freeze_panes = tl.cell(R0, T0)

# ---------------------------------------------------------------- Project Page
pp = wb.create_sheet('Project Page')
pp['A1'] = 'Project Page'
pp['A1'].font = TITLE
pp['A3'] = 'Select project ▸'
pp['A3'].font = BOLD
pp['B3'] = 'Example Project (slide)'
pp['B3'].font, pp['B3'].fill, pp['B3'].border = Font(name='Arial', size=11, bold=True, color='0000FF'), YELLOW, BOX
pp.merge_cells('B3:D3')
dvp = DataValidation(type='list', formula1=f'=Projects!$B${R0}:$B${R1}', allow_blank=True)
pp.add_data_validation(dvp)
dvp.add('B3')
pp['F3'] = '=IF(ISNA(MATCH(B3,Projects!$B$6:$B$55,0)),"◂ Pick a project from the dropdown","")'
pp['F3'].font = Font(name='Arial', size=10, color='C00000', bold=True)
pp['J1'] = 'row'
pp['J1'].font = GREY
pp['K1'] = f'=IFERROR(MATCH(B3,Projects!$B${R0}:$B${R1},0),0)'
pp['K1'].font = GREY
ROW = '$K$1'
fields = [('Engagement Manager', 'C', None), ('Starting Month', 'D', MON), ('Duration (months)', 'E', '0'),
          ('End Month', 'F', MON), ('Contracted Budget', 'G', MONEY), ('WIP to Date', 'H', MONEY),
          ('Budgeted (Contract − WIP)', 'I', MONEY), ('Remaining Months', 'J', '0'), ('Status', 'N', None)]
for i, (lab, col, fmt) in enumerate(fields):
    r = 5 + i
    pp.cell(r, 1, lab).font = BOLD
    c = pp.cell(r, 2, f'=IF({ROW}=0,"",INDEX(Projects!${col}${R0}:${col}${R1},{ROW})&"")' if fmt is None
                else f'=IF({ROW}=0,"",INDEX(Projects!${col}${R0}:${col}${R1},{ROW}))')
    c.font, c.border = GREEN, BOX
    if fmt:
        c.number_format = fmt
    pp.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
pp.column_dimensions['A'].width = 26
for col, w in zip('BCDEFGHI', [13, 13, 12, 14, 14, 14, 15, 15]):
    pp.column_dimensions[col].width = w
H = 16
header(pp, H, 1, ['Project Month', 'Calendar Month', 'Raw Weight', 'Plan Weight (Σ=1)', 'Original Plan (Contract × weight)',
                  'WIP Actual', 'Phased Budgeted', 'Total (WIP + Phased)', 'Actual / Forecast'])
pp.row_dimensions[H].height = 42
start = 'B6'
for k in range(1, MAXDUR + 1):
    r = H + k
    act = f'AND({ROW}>0,{k}<=N($B$7))'
    pp.cell(r, 1, f'=IF({act},{k},"")')
    pp.cell(r, 2, f'=IF(A{r}="","",EDATE($B$6,A{r}-1))').number_format = MON
    pp.cell(r, 3, f'=IF(A{r}="","",IF(INDEX(Phasing!$D${R0}:$D${R1},{ROW})=0,1,N(INDEX(Phasing!${L(W0)}${R0}:${L(W1)}${R1},{ROW},A{r}))))').number_format = '0.##'
    pp.cell(r, 4, f'=IF(A{r}="","",N(INDEX(Phasing!${L(NW0)}${R0}:${L(NW1)}${R1},{ROW},A{r})))').number_format = PCT
    pp.cell(r, 5, f'=IF(A{r}="","",D{r}*N($B$9))').number_format = MONEY
    col = mi(f'B{r}', TS)
    inrange = f'AND({col}>=1,{col}<={TL})'
    pp.cell(r, 6, f"=IF(A{r}=\"\",\"\",IF(AND({inrange},B{r}<={CM}),INDEX('WIP Log'!${L(M0)}${R0}:${L(M1)}${R1},{ROW},{col}),0))").number_format = MONEY
    pp.cell(r, 7, f"=IF(A{r}=\"\",\"\",IF({inrange},INDEX('Phased Budget'!${L(G0)}${R0}:${L(G1)}${R1},{ROW},{col}),0))").number_format = MONEY
    pp.cell(r, 8, f'=IF(A{r}="","",F{r}+G{r})').number_format = MONEY
    pp.cell(r, 9, f'=IF(A{r}="","",IF(B{r}<={CM},"Actual","Forecast"))')
    for c in range(1, 10):
        pp.cell(r, c).border = BOX
tr2 = H + MAXDUR + 1
pp.cell(tr2, 1, 'TOTAL').font = BOLD
for c in (4, 5, 6, 7, 8):
    x = pp.cell(tr2, c, f'=SUM({L(c)}{H + 1}:{L(c)}{H + MAXDUR})')
    x.font, x.number_format, x.border = BOLD, PCT if c == 4 else MONEY, Border(top=Side(style='medium'))
pp.cell(tr2 + 1, 1, 'Check: Total = Contract').font = GREY
pp.cell(tr2 + 1, 8, f'=IF({ROW}=0,"",IF(ABS(H{tr2}-N(B9))<0.005,"✓ ties","✗ diff "&TEXT(H{tr2}-N(B9),"#,##0")))').font = GREY
ch = BarChart()
ch.type, ch.grouping, ch.overlap = 'col', 'stacked', 100
ch.title = 'Monthly phasing – selected project'
ch.y_axis.title = 'Amount'
data = Reference(pp, min_col=6, max_col=7, min_row=H, max_row=H + MAXDUR)
ch.add_data(data, titles_from_data=True)
ch.set_categories(Reference(pp, min_col=2, min_row=H + 1, max_row=H + MAXDUR))
ln = LineChart()
ln.add_data(Reference(pp, min_col=5, min_row=H, max_row=H + MAXDUR), titles_from_data=True)
for se in ln.series:
    se.smooth = False
ch.gapWidth = 50
ch += ln
ch.height, ch.width = 9, 18
ch.x_axis.number_format = 'mmm yy'
ch.x_axis.delete = False
ch.y_axis.delete = False
pp.add_chart(ch, 'K4')

# ---------------------------------------------------------------- Aggregator
ag = wb.create_sheet('Aggregator')
title(ag, 'Aggregator', 'Σ WIP actuals (≤ Current Month) + Σ Phased Budgeted (after it) per month, against TRM. Enter TRM in the yellow column.')
header(ag, 5, 1, ['Month', 'Year', 'Σ WIP (actual)', 'Σ Phased Budgeted', 'Total (WIP + Phased)', 'TRM (input)', 'Variance (Total − TRM)'],
       [11, 7, 15, 15, 16, 14, 16])
A0, A1 = 6, 6 + TL - 1
for j in range(TL):
    r = A0 + j
    ag.cell(r, 1, f"='WIP Log'!{L(M0 + j)}5").number_format = 'mmm yy'
    ag.cell(r, 2, f'=YEAR(A{r})').number_format = '0'
    ag.cell(r, 3, f"=IF(A{r}<={CM},'WIP Log'!{L(M0 + j)}{tot},0)")
    ag.cell(r, 4, f"='Phased Budget'!{L(G0 + j)}{tot}")
    ag.cell(r, 5, f'=C{r}+D{r}')
    t = ag.cell(r, 6)
    t.font, t.fill = BLUE, YELLOW
    ag.cell(r, 7, f'=IF(F{r}="","",E{r}-F{r})')
    for c in range(1, 8):
        ag.cell(r, c).border = BOX
        if c >= 3:
            ag.cell(r, c).number_format = MONEY
    ag.cell(r, 3).font = ag.cell(r, 4).font = GREEN
ag['F5'].comment = Comment('TRM: monthly target, entered by hand (per user). Leave it blank where you have no target.', 'Claude')
ag.conditional_formatting.add(f'A{A0}:G{A1}', FormulaRule(formula=[f'AND(YEAR($A{A0})=YEAR({CM}),MONTH($A{A0})=MONTH({CM}))'],
                              fill=PatternFill('solid', fgColor='FCE4D6'), font=Font(bold=True)))
tr3 = A1 + 1
ag.cell(tr3, 1, 'TOTAL').font = BOLD
for c in range(3, 8):
    x = ag.cell(tr3, c, f'=SUM({L(c)}{A0}:{L(c)}{A1})')
    x.font, x.number_format, x.border = BOLD, MONEY, Border(top=Side(style='medium'))
# annual table
ag['I5'] = 'Annual (calendar year)'
ag['I5'].font = BOLD
header(ag, 6, 9, ['Year', 'Σ WIP (actual)', 'Σ Phased Budget', 'WIP + Phased Budget', 'TRM', 'Variance'],
       [8, 15, 15, 17, 14, 14])
for i in range(3):
    r = 7 + i
    ag.cell(r, 9, f'=YEAR({TS})+{i}').number_format = '0'
    for c, src in zip((10, 11, 13), ('C', 'D', 'F')):
        ag.cell(r, c, f'=SUMIFS({src}${A0}:{src}${A1},$B${A0}:$B${A1},$I{r})')
    ag.cell(r, 12, f'=J{r}+K{r}')
    ag.cell(r, 14, f'=L{r}-M{r}')
    for c in range(9, 15):
        ag.cell(r, c).border = BOX
        if c > 9:
            ag.cell(r, c).number_format = MONEY
ag.cell(10, 9, 'Total').font = BOLD
for c in range(10, 15):
    x = ag.cell(10, c, f'=SUM({L(c)}7:{L(c)}9)')
    x.font, x.number_format = BOLD, MONEY
# reconciliation
ag['I12'] = 'Reconciliation'
ag['I12'].font = BOLD
recon = [('Σ Contracted Budget', f'=Projects!G{R1 + 2}'),
         ('Σ WIP to date (all projects)', f'=Projects!H{R1 + 2}'),
         ('Σ Phased Budgeted (on timeline)', f'=D{tr3}'),
         ('Difference (unphased / off-timeline)', '=L13-L14-L15')]
for i, (lab, f) in enumerate(recon):
    r = 13 + i
    ag.cell(r, 9, lab)
    ag.merge_cells(start_row=r, start_column=9, end_row=r, end_column=11)
    c = ag.cell(r, 12, f)
    c.number_format, c.border = MONEY, BOX
ag['M16'] = '=IF(ABS(L16)<0.005,"✓ fully phased","Check Projects → Status")'
ag['M16'].font = Font(name='Arial', size=10, bold=True, color='C0392B')
bc = BarChart()
bc.type, bc.grouping, bc.overlap = 'col', 'stacked', 100
bc.title = 'Σ WIP + Σ Phased Budgeted vs TRM per month'
bc.add_data(Reference(ag, min_col=3, max_col=4, min_row=5, max_row=A1), titles_from_data=True)
bc.set_categories(Reference(ag, min_col=1, min_row=A0, max_row=A1))
lc = LineChart()
lc.add_data(Reference(ag, min_col=6, min_row=5, max_row=A1), titles_from_data=True)
for se in lc.series:
    se.smooth = False
bc.gapWidth = 50
bc += lc
bc.height, bc.width = 9, 24
bc.x_axis.number_format = 'mmm yy'
bc.x_axis.delete = False
bc.y_axis.delete = False
ag.add_chart(bc, 'I19')
ag.freeze_panes = 'A6'

# sheet order: README, Projects, Phasing, WIP Log, Project Page, Timeline, Aggregator, Phased Budget
order = ['README', 'Projects', 'Phasing', 'WIP Log', 'Project Page', 'Timeline', 'Aggregator', 'Phased Budget']
wb._sheets = [wb[n] for n in order]
for ws, color in zip(wb._sheets, ['808080', 'C0392B', 'F39C12', 'F39C12', '2E86C1', '2E86C1', '2E86C1', '808080']):
    ws.sheet_properties.tabColor = color
wb.active = 1
wb.save(OUT)
print('saved', OUT)
