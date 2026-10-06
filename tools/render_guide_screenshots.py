"""Build a demo copy of the workbook and render one screenshot per sheet area."""
import os, shutil, subprocess, sys
from datetime import date
from openpyxl import load_workbook

SRC, OUT = sys.argv[1], sys.argv[2]
RECALC = sys.argv[3]
os.makedirs(OUT, exist_ok=True)
demo = os.path.join(OUT, 'demo.xlsx')
shutil.copy(SRC, demo)

wb = load_workbook(demo)
p, ph, wl, ag = wb['Projects'], wb['Phasing'], wb['WIP Log'], wb['Aggregator']


def add(r, name, em, start, dur, contract, weights=None, wip=None):
    p.cell(r, 2, name); p.cell(r, 3, em); p.cell(r, 4, start); p.cell(r, 5, dur); p.cell(r, 7, contract)
    for i, w in enumerate(weights or []):
        ph.cell(r, 9 + i, w)
    for col, v in (wip or {}).items():   # col = 0-based timeline index (0 = Jan 2026)
        wl.cell(r, 4 + col, v)


add(7, 'Data Platform Migration', 'M. Laurent', date(2026, 6, 1), 12, 120000, None, {5: 9000, 6: 9000, 7: 9000, 8: 9000, 9: 9000})
add(8, 'CRM Rollout', 'A. Okafor', date(2026, 1, 1), 9, 45000, [1, 1, 2, 2, 2, 1, 1, 1, 1], {i: 5000 for i in range(9)})
add(9, 'Security Audit', 'S. Novak', date(2027, 3, 1), 4, 30000, [1, 2, 2, 1])
add(10, 'Cloud Cost Review', 'R. Patel', date(2026, 8, 1), 6, 18000, [2, 2, 1, 1, 1, 1], {7: 4000, 8: 4000, 9: 3000})
add(11, 'Analytics Pilot', 'L. Chen', date(2026, 9, 1), 3, 9000, None, {8: 6000, 9: 4500})
for j in range(24):
    ag.cell(6 + j, 6, 25000)
wb.save(demo)
print(subprocess.run([sys.executable, RECALC, demo, '120'], capture_output=True, text=True).stdout[:120])

from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
EMPTY = list(range(14, 56))   # unused project rows 14-55 (keeps 6 sample rows + 2 blank + totals row 57)


def cols(a, b):
    return [L(i) for i in range(CI(a), CI(b) + 1)]


# key: (sheet, print area, rows to hide, columns to hide)
SHOTS = {
    'readme': ('README', 'B1:C38', [], []),
    'projectpage': ('Project Page', 'A1:I42', [], []),
    'projectchart': ('Project Page', 'K3:X22', [], []),
    'aggregator': ('Aggregator', 'A1:G42', [], []),
    'aggannual': ('Aggregator', 'I5:X37', [], []),
}


def chunk(key, sheet, last_row, keep, hrows):
    """Screenshot only the column ranges in `keep` (others hidden); print area spans them all."""
    idx = set()
    for a_, b_ in keep:
        idx.update(range(CI(a_), CI(b_) + 1))
    hi = max(idx)
    SHOTS[key] = (sheet, f'A1:{L(hi)}{last_row}', hrows, [L(i) for i in range(1, hi + 1) if i not in idx])


# Wide sheets in portrait-width pieces (12 months / 12 weights each), names repeated on every piece
chunk('projects1', 'Projects', 57, [('A', 'G')], EMPTY)
chunk('projects2', 'Projects', 57, [('A', 'B'), ('H', 'N')], EMPTY)
chunk('phasing0', 'Phasing', 57, [('A', 'H')], EMPTY)
chunk('phasing1', 'Phasing', 57, [('A', 'B'), ('I', 'T')], EMPTY)
chunk('phasing2', 'Phasing', 57, [('A', 'B'), ('U', 'AF')], EMPTY)
chunk('phasing3', 'Phasing', 57, [('A', 'B'), ('AH', 'AS')], EMPTY)
chunk('phasing4', 'Phasing', 57, [('A', 'B'), ('AT', 'BF')], EMPTY)
for n, (a_, b_) in enumerate([('D', 'O'), ('P', 'AA'), ('AB', 'AM')], 1):
    chunk(f'wiplog{n}', 'WIP Log', 57, [('A', 'C'), (a_, b_)], EMPTY)
for n, (a_, b_) in enumerate([('F', 'Q'), ('R', 'AC'), ('AD', 'AO')], 1):
    chunk(f'timeline{n}', 'Timeline', 13, [('A', 'E'), (a_, b_)], [])
chunk('phased0', 'Phased Budget', 57, [('A', 'G')], EMPTY)
for n, (a_, b_) in enumerate([('H', 'S'), ('T', 'AE'), ('AF', 'AQ')], 1):
    chunk(f'phased{n}', 'Phased Budget', 57, [('A', 'B'), (a_, b_)], EMPTY)
for key, (sheet, area, hrows, hcols) in SHOTS.items():
    wb = load_workbook(demo)
    for ws in wb.worksheets:
        ws.sheet_state = 'visible' if ws.title == sheet else 'hidden'
    ws = wb[sheet]
    wb.active = wb.worksheets.index(ws)
    ws.print_area = area
    for r in hrows:
        ws.row_dimensions[r].hidden = True
    # openpyxl merges identical adjacent columns into one <col min max> group; split them so we can hide single columns
    for k, d in list(ws.column_dimensions.items()):
        if d.max and d.min and d.max > d.min:
            lo, hi, w = d.min, d.max, d.width
            d.max = lo
            for i in range(lo + 1, hi + 1):
                nd = ws.column_dimensions[L(i)]
                nd.min = nd.max = i
                nd.width = w
    for c in hcols:
        ws.column_dimensions[c].hidden = True
    ws.freeze_panes = None
    # Wrap long notes in the top rows inside the visible columns so they are not clipped at the screenshot edge
    import math
    from openpyxl.styles import Alignment
    from openpyxl.utils.cell import range_boundaries
    c0, r0, c1, r1 = range_boundaries(area)
    hidden = set(CI(c) for c in hcols)
    for row in range(r0, min(r0 + 4, r1 + 1)):
        for col in range(c0, c1 + 1):
            cell = ws.cell(row, col)
            if col in hidden or not isinstance(cell.value, str) or cell.value.startswith('=') or len(cell.value) < 30:
                continue
            end = col
            while end + 1 <= c1 and ws.cell(row, end + 1).value in (None, ''):
                end += 1
            width = sum((ws.column_dimensions[L(i)].width or 8.43) for i in range(col, end + 1) if i not in hidden)
            if end > col:
                ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=end)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            lines = max(1, math.ceil(len(cell.value) * 1.05 / max(width, 1)))
            ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 13, 13 * lines + 2)
    ws.print_title_rows = None
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.gridLines = False
    ws.page_margins.left = ws.page_margins.right = ws.page_margins.top = ws.page_margins.bottom = 0.2
    ws.oddHeader.center.text = ws.oddFooter.center.text = ''
    x = os.path.join(OUT, f'{key}.xlsx')
    wb.save(x)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', OUT, x], capture_output=True)
    subprocess.run(['pdftoppm', '-png', '-r', '170', '-f', '1', '-l', '1', '-singlefile',
                    os.path.join(OUT, f'{key}.pdf'), os.path.join(OUT, key)], capture_output=True)
    # trim whitespace
    from PIL import Image, ImageChops
    im = Image.open(os.path.join(OUT, f'{key}.png')).convert('RGB')
    bg = Image.new('RGB', im.size, (255, 255, 255))
    bbox = ImageChops.difference(im, bg).getbbox()
    if bbox:
        im = im.crop((max(0, bbox[0] - 8), max(0, bbox[1] - 8), min(im.width, bbox[2] + 8), min(im.height, bbox[3] + 8)))
    im.save(os.path.join(OUT, f'{key}.png'))
    print(key, im.size)
